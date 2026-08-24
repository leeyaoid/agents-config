#!/usr/bin/env python3
"""Validate required styles and reject direct typography in generated DOCX files."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from generate_official_doc import BUILTIN_PARAGRAPH_STYLE_IDS, CHARACTER_STYLES, PARAGRAPH_STYLES


ALLOWED_RUN_PROPERTY_TAGS = {qn("w:rStyle")}


def _iter_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def validate_style_usage(path: str | Path):
    document = Document(path)
    errors = []
    warnings = []
    style_names = {style.name for style in document.styles}
    required = set(PARAGRAPH_STYLES.values()) | set(CHARACTER_STYLES.values())
    missing = sorted(required - style_names)
    if missing:
        errors.append("缺少必需样式: " + "、".join(missing))

    for key, expected_style_id in BUILTIN_PARAGRAPH_STYLE_IDS.items():
        style_name = PARAGRAPH_STYLES[key]
        try:
            actual_style_id = document.styles[style_name].style_id
        except KeyError:
            continue
        if actual_style_id != expected_style_id:
            errors.append(f"{style_name}未复用Word内置样式: styleId={actual_style_id}，应为{expected_style_id}")

    expected_paragraph_styles = set(PARAGRAPH_STYLES.values())
    for index, paragraph in enumerate(_iter_paragraphs(document), 1):
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name not in expected_paragraph_styles:
            errors.append(f"第{index}个非空段落未使用规定段落样式: {style_name or '无样式'} | {text[:24]}")
        for run_index, run in enumerate(paragraph.runs, 1):
            rpr = run._r.rPr
            if rpr is None:
                continue
            direct_tags = {child.tag for child in rpr if child.tag not in ALLOWED_RUN_PROPERTY_TAGS}
            if direct_tags:
                readable = ", ".join(sorted(tag.split("}")[-1] for tag in direct_tags))
                errors.append(f"段落{index}运行{run_index}含直接字符格式: {readable} | {text[:24]}")
        if re.search(r"[\u3400-\u9fff]", text) and ('"' in text or "'" in text):
            errors.append(f"第{index}个中文段落仍含英文直引号: {text[:24]}")

    if not document.paragraphs:
        warnings.append("正文中没有段落")
    if any(style.name.startswith("公文-") for style in document.styles):
        errors.append("仍存在带‘公文-’前缀的样式")
    if any("中文引号" in style.name for style in document.styles):
        errors.append("仍存在独立引号样式")
    return {"passed": not errors, "errors": errors, "warnings": warnings}


def main():
    parser = argparse.ArgumentParser(description="校验DOCX是否真正使用Word样式")
    parser.add_argument("docx")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    result = validate_style_usage(args.docx)
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print("Style validation: " + ("PASSED" if result["passed"] else "FAILED"))
        for message in result["errors"]:
            print("ERROR: " + message)
        for message in result["warnings"]:
            print("WARNING: " + message)
    sys.exit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()

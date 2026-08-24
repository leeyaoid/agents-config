#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate GB/T 9704 DOCX files through Word styles, not direct formatting."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any, Dict, Mapping

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


DEFAULT_TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "assets" / "official_document_normal_template.docx"

PARAGRAPH_STYLES = {
    "normal": "Normal",
    "issuer_mark": "发文机关标志",
    "classification": "密级与紧急程度",
    "document_number": "发文字号",
    "red_separator": "红色分隔线",
    "title": "Title",
    "recipient": "主送机关",
    "heading1": "Heading 1",
    "heading2": "Heading 2",
    "heading3": "Heading 3",
    "heading4": "Heading 4",
    "attachment": "附件说明",
    "signature": "发文机关署名",
    "date": "成文日期",
    "signer_line": "签发人行",
    "note": "附注",
    "imprint": "版记",
    "imprint_separator": "版记分隔线",
    "page_number": "页码段落",
}

CHARACTER_STYLES = {
    "signer_label": "签发人标签",
    "signer_name": "签发人姓名",
    "page_number": "页码字符",
}

BUILTIN_PARAGRAPH_STYLE_IDS = {
    "normal": "Normal",
    "title": "Title",
    "heading1": "Heading1",
    "heading2": "Heading2",
    "heading3": "Heading3",
    "heading4": "Heading4",
}

STYLE_ALIASES = {
    "正文": "normal",
    "一级标题": "heading1",
    "标题 1": "heading1",
    "二级标题": "heading2",
    "标题 2": "heading2",
    "三级标题": "heading3",
    "标题 3": "heading3",
    "四级标题": "heading4",
    "标题 4": "heading4",
    "附件": "attachment",
    "附注": "note",
}

CLOSING_PHRASES = {
    "通知": "特此通知。",
    "报告": "特此报告。",
    "请示": "妥否，请批示。",
    "函": "请予研究函复。",
    "通报": "特此通报。",
}


def _ensure_style(document: Document, name: str, style_type: WD_STYLE_TYPE):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, style_type)


def _clear_style_formatting(style):
    for tag in ("w:pPr", "w:rPr"):
        properties = style.element.find(qn(tag))
        if properties is not None:
            for child in list(properties):
                properties.remove(child)


def _set_style_font(
    style,
    east_asia: str,
    size_pt: float,
    *,
    bold: bool = False,
    color: str = "000000",
    latin_font: str = "Times New Roman",
    east_asia_hint: bool = False,
):
    style.font.name = latin_font
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin_font)
    rfonts.set(qn("w:hAnsi"), latin_font)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin_font)
    if east_asia_hint:
        rfonts.set(qn("w:hint"), "eastAsia")


def _set_outline_level(style, level: int):
    ppr = style.element.get_or_add_pPr()
    old = ppr.find(qn("w:outlineLvl"))
    if old is not None:
        ppr.remove(old)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    ppr.append(outline)


def _set_style_border(style, *, color: str, size: str):
    ppr = style.element.get_or_add_pPr()
    old = ppr.find(qn("w:pBdr"))
    if old is not None:
        ppr.remove(old)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    ppr.append(borders)


def _configure_paragraph_style(
    document: Document,
    name: str,
    *,
    font: str,
    size: float,
    bold: bool = False,
    color: str = "000000",
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
    first_line_pt: float = 0,
    left_pt: float = 0,
    right_pt: float = 0,
    line_pt: float = 28,
    before_pt: float = 0,
    after_pt: float = 0,
    keep_next: bool = False,
    outline_level: int | None = None,
):
    style = _ensure_style(document, name, WD_STYLE_TYPE.PARAGRAPH)
    _clear_style_formatting(style)
    _set_style_font(style, font, size, bold=bold, color=color, east_asia_hint=True)
    style.quick_style = True
    fmt = style.paragraph_format
    fmt.alignment = alignment
    fmt.first_line_indent = Pt(first_line_pt)
    fmt.left_indent = Pt(left_pt)
    fmt.right_indent = Pt(right_pt)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(line_pt)
    fmt.space_before = Pt(before_pt)
    fmt.space_after = Pt(after_pt)
    fmt.keep_with_next = keep_next
    fmt.widow_control = True
    if outline_level is not None:
        _set_outline_level(style, outline_level)
    return style


def create_official_styles(document: Document) -> Dict[str, str]:
    _configure_paragraph_style(document, PARAGRAPH_STYLES["normal"], font="仿宋", size=16, first_line_pt=32)
    _configure_paragraph_style(document, PARAGRAPH_STYLES["issuer_mark"], font="方正小标宋_GBK", size=26, color="FF0000", alignment=WD_ALIGN_PARAGRAPH.CENTER, line_pt=32, after_pt=12)
    _configure_paragraph_style(document, PARAGRAPH_STYLES["classification"], font="黑体", size=16)
    _configure_paragraph_style(document, PARAGRAPH_STYLES["document_number"], font="仿宋", size=16, alignment=WD_ALIGN_PARAGRAPH.CENTER, before_pt=12, after_pt=4)
    red_line = _configure_paragraph_style(document, PARAGRAPH_STYLES["red_separator"], font="宋体", size=1, line_pt=4, after_pt=24)
    _set_style_border(red_line, color="FF0000", size="18")
    _configure_paragraph_style(document, PARAGRAPH_STYLES["title"], font="方正小标宋_GBK", size=22, alignment=WD_ALIGN_PARAGRAPH.CENTER, after_pt=28, keep_next=True, outline_level=0)
    _configure_paragraph_style(document, PARAGRAPH_STYLES["recipient"], font="仿宋", size=16, keep_next=True)
    _configure_paragraph_style(document, PARAGRAPH_STYLES["heading1"], font="黑体", size=16, first_line_pt=32, keep_next=True, outline_level=0)
    _configure_paragraph_style(document, PARAGRAPH_STYLES["heading2"], font="楷体", size=16, first_line_pt=32, keep_next=True, outline_level=1)
    _configure_paragraph_style(document, PARAGRAPH_STYLES["heading3"], font="仿宋", size=16, bold=True, first_line_pt=32, keep_next=True, outline_level=2)
    _configure_paragraph_style(document, PARAGRAPH_STYLES["heading4"], font="仿宋", size=16, first_line_pt=32, keep_next=True, outline_level=3)
    _configure_paragraph_style(document, PARAGRAPH_STYLES["attachment"], font="仿宋", size=16, first_line_pt=32, before_pt=28)
    _configure_paragraph_style(document, PARAGRAPH_STYLES["signature"], font="仿宋", size=16, alignment=WD_ALIGN_PARAGRAPH.RIGHT, right_pt=64, before_pt=28)
    _configure_paragraph_style(document, PARAGRAPH_STYLES["date"], font="仿宋", size=16, alignment=WD_ALIGN_PARAGRAPH.RIGHT, right_pt=64)
    _configure_paragraph_style(document, PARAGRAPH_STYLES["signer_line"], font="仿宋", size=16, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    _configure_paragraph_style(document, PARAGRAPH_STYLES["note"], font="仿宋", size=16, first_line_pt=32)
    imprint = _configure_paragraph_style(document, PARAGRAPH_STYLES["imprint"], font="仿宋", size=14, line_pt=20)
    imprint.paragraph_format.tab_stops.add_tab_stop(Cm(15.6), WD_TAB_ALIGNMENT.RIGHT)
    imprint_line = _configure_paragraph_style(document, PARAGRAPH_STYLES["imprint_separator"], font="宋体", size=1, line_pt=4)
    _set_style_border(imprint_line, color="000000", size="8")
    _configure_paragraph_style(document, PARAGRAPH_STYLES["page_number"], font="宋体", size=14, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_pt=14)

    for key, font, size in (
        ("signer_label", "仿宋", 16),
        ("signer_name", "楷体", 16),
        ("page_number", "宋体", 14),
    ):
        style = _ensure_style(document, CHARACTER_STYLES[key], WD_STYLE_TYPE.CHARACTER)
        _clear_style_formatting(style)
        _set_style_font(style, font, size, east_asia_hint=True)
        style.quick_style = True
    return dict(PARAGRAPH_STYLES)


def _configure_page(document: Document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.7)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    section.footer_distance = Cm(1.75)
    sect_pr = section._sectPr
    doc_grid = sect_pr.find(qn("w:docGrid"))
    if doc_grid is None:
        doc_grid = OxmlElement("w:docGrid")
        sect_pr.append(doc_grid)
    doc_grid.set(qn("w:type"), "linesAndChars")
    doc_grid.set(qn("w:linePitch"), "560")
    doc_grid.set(qn("w:charSpace"), "0")


def _normalize_chinese_quotes(text: str) -> str:
    if not re.search(r"[\u3400-\u9fff]", text):
        return text
    text = re.sub(r'"([^"\n]+)"', r'“\1”', text)
    return re.sub(r"'([^'\n]+)'", r"‘\1’", text)


def _add_styled_paragraph(document: Document, text: str, style_key: str):
    paragraph = document.add_paragraph(style=PARAGRAPH_STYLES[style_key])
    if text:
        paragraph.add_run(_normalize_chinese_quotes(text))
    return paragraph


def _add_page_number(document: Document):
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.style = PARAGRAPH_STYLES["page_number"]
    left = paragraph.add_run("— ")
    left.style = CHARACTER_STYLES["page_number"]
    for tag, value in (("w:fldChar", "begin"), ("w:instrText", " PAGE "), ("w:fldChar", "separate"), ("w:t", "1"), ("w:fldChar", "end")):
        run = paragraph.add_run()
        run.style = CHARACTER_STYLES["page_number"]
        element = OxmlElement(tag)
        if tag == "w:fldChar":
            element.set(qn("w:fldCharType"), value)
        else:
            element.text = value
            if tag == "w:instrText":
                element.set(qn("xml:space"), "preserve")
        run._r.append(element)
    right = paragraph.add_run(" —")
    right.style = CHARACTER_STYLES["page_number"]


def _detect_body_style(text: str) -> str:
    if re.match(r"^[一二三四五六七八九十百]+、", text):
        return "heading1"
    if re.match(r"^（[一二三四五六七八九十百]+）", text):
        return "heading2"
    if re.match(r"^\d+[\.．]", text):
        return "heading3"
    if re.match(r"^（\d+）", text):
        return "heading4"
    return "normal"


def _body_item(item: Any) -> tuple[str, str]:
    if isinstance(item, str):
        return item, _detect_body_style(item)
    if isinstance(item, Mapping):
        text = str(item.get("text", ""))
        requested = str(item.get("style", "正文"))
        key = STYLE_ALIASES.get(requested, requested)
        if key not in PARAGRAPH_STYLES:
            raise ValueError(f"未知正文样式: {requested}")
        return text, key
    raise TypeError("body 项必须是字符串或包含 text/style 的字典")


def _add_signer(document: Document, name: str):
    paragraph = document.add_paragraph(style=PARAGRAPH_STYLES["signer_line"])
    label = paragraph.add_run("签发人：")
    label.style = CHARACTER_STYLES["signer_label"]
    value = paragraph.add_run(name)
    value.style = CHARACTER_STYLES["signer_name"]


def _add_imprint(document: Document, content: Mapping[str, Any]):
    copy_to = content.get("copy_to")
    print_org = content.get("print_org") or content.get("issuer_office")
    print_date = content.get("print_date") or content.get("issue_date")
    if not any((copy_to, print_org, print_date)):
        return
    _add_styled_paragraph(document, "", "imprint_separator")
    if copy_to:
        _add_styled_paragraph(document, f"抄送：{copy_to}。", "imprint")
    if print_org or print_date:
        _add_styled_paragraph(document, f"{print_org or ''}\t{print_date or ''}印发", "imprint")
    _add_styled_paragraph(document, "", "imprint_separator")


def create_official_document(doc_type: str, content: Mapping[str, Any], template_path: str | Path | None = None) -> Document:
    template = Path(template_path) if template_path else DEFAULT_TEMPLATE_PATH
    document = Document(str(template)) if template.exists() else Document()
    _configure_page(document)
    create_official_styles(document)

    include_header = bool(content.get("include_header", True))
    issuer = str(content.get("issuer", ""))
    if content.get("classification"):
        classification = str(content["classification"])
        if content.get("classification_period"):
            classification += "★" + str(content["classification_period"])
        _add_styled_paragraph(document, classification, "classification")
    if content.get("urgency"):
        _add_styled_paragraph(document, str(content["urgency"]), "classification")
    if include_header and issuer:
        _add_styled_paragraph(document, issuer if issuer.endswith("文件") else issuer + "文件", "issuer_mark")
    if include_header and content.get("doc_number"):
        _add_styled_paragraph(document, str(content["doc_number"]), "document_number")
    if content.get("signer"):
        _add_signer(document, str(content["signer"]))
    if include_header:
        _add_styled_paragraph(document, "", "red_separator")
    if content.get("title"):
        _add_styled_paragraph(document, str(content["title"]), "title")
    if content.get("recipient"):
        _add_styled_paragraph(document, str(content["recipient"]).rstrip("：:") + "：", "recipient")

    for item in content.get("body", []):
        text, style_key = _body_item(item)
        _add_styled_paragraph(document, text, style_key)

    closing = content.get("closing", CLOSING_PHRASES.get(doc_type))
    if closing:
        _add_styled_paragraph(document, str(closing), "normal")

    attachments = content.get("attachments")
    if not attachments and content.get("attachment"):
        attachments = [content["attachment"]]
    if attachments:
        values = list(attachments) if isinstance(attachments, (list, tuple)) else [attachments]
        if len(values) == 1:
            _add_styled_paragraph(document, f"附件：{values[0]}", "attachment")
        else:
            _add_styled_paragraph(document, "附件：", "attachment")
            for index, value in enumerate(values, 1):
                _add_styled_paragraph(document, f"{index}. {value}", "attachment")

    signature = content.get("issuer_signature", issuer if include_header else "")
    if signature:
        _add_styled_paragraph(document, str(signature), "signature")
    if content.get("date"):
        _add_styled_paragraph(document, str(content["date"]), "date")
    if content.get("note"):
        note = str(content["note"])
        if not (note.startswith("（") and note.endswith("）")):
            note = f"（{note}）"
        _add_styled_paragraph(document, note, "note")

    _add_imprint(document, content)
    if content.get("page_numbers", True):
        _add_page_number(document)
    document.core_properties.title = str(content.get("title", ""))
    document.core_properties.author = issuer
    document.core_properties.subject = f"{doc_type}类公文"
    return document


def generate_document(doc_type: str, content: Mapping[str, Any], output_path: str | Path, template_path: str | Path | None = None) -> str:
    document = create_official_document(doc_type, content, template_path)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return str(output)


def _demo_content() -> Dict[str, Any]:
    return {
        "issuer": "示例公司",
        "doc_number": "示例〔2026〕1号",
        "title": "关于年度重点工作的报告",
        "recipient": "上级单位",
        "body": ["现将年度重点工作情况报告如下。", "一、总体情况", "年度重点任务有序推进。", "（一）主要成效", "项目建设取得阶段性进展。", "1. 完善工作机制", "进一步健全统筹协调机制。", "（1）细化责任分工", "各项任务均已落实责任部门。"],
        "issuer_signature": "示例公司",
        "date": "2026年6月29日",
        "copy_to": "公司各部门",
        "print_org": "示例公司办公室",
        "print_date": "2026年6月29日",
    }


def main():
    parser = argparse.ArgumentParser(description="生成使用Word样式的公文DOCX")
    parser.add_argument("--type", default="报告")
    parser.add_argument("--input-json")
    parser.add_argument("--template", help="Word基础模板；默认使用skill内置Normal模板")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    content = json.loads(Path(args.input_json).read_text(encoding="utf-8")) if args.input_json else _demo_content()
    print(generate_document(args.type, content, args.output, args.template))


if __name__ == "__main__":
    main()
